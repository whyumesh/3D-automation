import pandas as pd
import os
from jinja2 import Environment, FileSystemLoader
import win32com.client as win32
from datetime import datetime as dt
from openpyxl import load_workbook
import glob

# Load Jinja2 Template for email
env = Environment(loader=FileSystemLoader('.'))
template = env.get_template("email_template_ZBM.html")

# Get current date
z = dt.today()
current_date = z.date()

# Find the most recent ZBM folders
def find_latest_folder(pattern):
    """Find the most recently created folder matching the pattern"""
    folders = glob.glob(pattern)
    if not folders:
        return None
    return max(folders, key=os.path.getctime)

# Get the current directory
current_dir = os.path.dirname(os.path.abspath(__file__))
if not current_dir:
    current_dir = os.getcwd()

print(f"📂 Working directory: {current_dir}")

# Locate the generated folders
consolidated_folder = find_latest_folder(os.path.join(current_dir, "ZBM_Consolidated_Files_*"))
reports_folder = find_latest_folder(os.path.join(current_dir, "ZBM_Reports_*"))

if not consolidated_folder:
    print("❌ Error: Could not find ZBM_Consolidated_Files folder. Please run create_zbm_consolidated_files.py first.")
    exit(1)

if not reports_folder:
    print("❌ Error: Could not find ZBM_Reports folder. Please run create_zbm_hierarchical_reports.py first.")
    exit(1)

print(f"✅ Found consolidated files folder: {consolidated_folder}")
print(f"✅ Found reports folder: {reports_folder}")

# Debug: List files in consolidated folder
print(f"\n🔍 Files in consolidated folder:")
if consolidated_folder and os.path.exists(consolidated_folder):
    files = os.listdir(consolidated_folder)
    print(f"   Total files: {len(files)}")
    for f in files[:5]:  # Show first 5 files
        print(f"   - {f}")
else:
    print("   Folder not accessible!")

# Debug: List files in reports folder
print(f"\n🔍 Files in reports folder:")
if reports_folder and os.path.exists(reports_folder):
    files = os.listdir(reports_folder)
    print(f"   Total files: {len(files)}")
    for f in files[:5]:  # Show first 5 files
        print(f"   - {f}")
else:
    print("   Folder not accessible!")

# Read Sample Master Tracker to get ZBM details
print("📖 Reading Sample Master Tracker.xlsx...")
df = pd.read_excel('Sample Master Tracker.xlsx')

# Filter for ZBM codes starting with "ZN"
df = df[df['ZBM Terr Code'].astype(str).str.startswith('ZN')]

# Get unique ZBMs with their details
zbms = df[['ZBM Terr Code', 'ZBM Name', 'ZBM EMAIL_ID']].drop_duplicates().sort_values('ZBM Terr Code')
print(f"📋 Found {len(zbms)} unique ZBMs to process")

# Initialize Outlook
outlook = win32.Dispatch("Outlook.Application")

# Create output directory for sent email logs
output_dir = os.path.dirname(os.path.abspath(__file__))
email_log_folder = os.path.join(output_dir, f'ZBM_Email_Logs_{current_date}')
os.makedirs(email_log_folder, exist_ok=True)

def read_summary_report(zbm_code, zbm_name):
    """Read the summary report Excel file for a ZBM and extract data as HTML table"""
    try:
        # Find the summary report file - use wildcard pattern without safe_zbm_name
        pattern = os.path.join(reports_folder, f"ZBM_Summary_{zbm_code}_*.xlsx")
        files = glob.glob(pattern)
        
        if not files:
            print(f"   ⚠️ Warning: No summary report found for {zbm_code}")
            print(f"      Searched pattern: {pattern}")
            return None
        
        report_file = os.path.abspath(files[0])
        
        # Verify file exists
        if not os.path.exists(report_file):
            print(f"   ❌ Summary report file does not exist: {report_file}")
            return None
        
        print(f"   📊 Reading summary report: {os.path.basename(report_file)}")
        
        # Read the Excel file
        wb = load_workbook(report_file)
        ws = wb['ZBM']
        
        # Find header row and starting column (looking for "Area Name")
        header_row = None
        start_col = None
        for row_idx in range(1, 15):
            for col_idx in range(1, 20):
                cell_value = ws.cell(row=row_idx, column=col_idx).value
                if cell_value and 'Area Name' in str(cell_value):
                    header_row = row_idx
                    start_col = col_idx
                    break
            if header_row:
                break

        if not header_row or not start_col:
            print(f"   ⚠️ Warning: Could not find header row in summary report")
            return None

        # Read headers starting from start_col
        headers = []
        for col_idx in range(start_col, ws.max_column + 1):
            header_val = ws.cell(row=header_row, column=col_idx).value
            if header_val is None or str(header_val).strip() == "":
                break
            headers.append(str(header_val).strip())

        # Read data rows
        data = []
        empty_row_count = 0
        for row_idx in range(header_row + 1, ws.max_row + 1):
            row_data = []
            has_any_value = False

            for col_offset in range(len(headers)):
                col_idx = start_col + col_offset
                cell_value = ws.cell(row=row_idx, column=col_idx).value
                
                # Check if cell has any meaningful value
                if cell_value is not None and str(cell_value).strip() != "":
                    has_any_value = True
                
                row_data.append(cell_value)

            # If row has at least one value, add it to data
            if has_any_value:
                data.append(row_data)
                empty_row_count = 0
            else:
                empty_row_count += 1
                # Stop only after 2 consecutive completely empty rows
                if empty_row_count >= 2:
                    break
        
        # Create DataFrame
        df_summary = pd.DataFrame(data, columns=headers)
        
        # Remove rows where the first column is "None", empty, or contains only whitespace
        if len(df_summary) > 0 and len(df_summary.columns) > 0:
            first_col = df_summary.columns[0]
            df_summary = df_summary[
                (df_summary[first_col].notna()) & 
                (df_summary[first_col].astype(str).str.strip() != "") &
                (df_summary[first_col].astype(str).str.strip().str.lower() != "none")
            ]
        
        # Reset index after filtering
        df_summary = df_summary.reset_index(drop=True)
        
        # Convert to HTML table with styling
        html_table = df_summary.to_html(index=False, border=1, classes='summary-table', na_rep='-')
        
        return html_table
        
    except Exception as e:
        print(f"   ❌ Error reading summary report for {zbm_code}: {e}")
        import traceback
        traceback.print_exc()
        return None

def get_abm_emails_for_zbm(zbm_code):
    """Get all ABM email addresses under a specific ZBM for CC"""
    zbm_data = df[df['ZBM Terr Code'] == zbm_code]
    abm_emails = zbm_data['ABM EMAIL_ID'].dropna().unique()
    
    # Filter out invalid emails (0, '0', empty strings)
    abm_emails = [email for email in abm_emails if email and str(email) not in ['0', '0.0']]
    
    return '; '.join(abm_emails)

# Process each ZBM and send emails
email_count = 0
for _, zbm_row in zbms.iterrows():
    zbm_code = zbm_row['ZBM Terr Code']
    zbm_name = zbm_row['ZBM Name']
    zbm_email = zbm_row['ZBM EMAIL_ID']
    
    print(f"\n🔄 Processing ZBM: {zbm_code} - {zbm_name}")
    
    # Skip if no valid email
    if not zbm_email or str(zbm_email) in ['0', '0.0']:
        print(f"   ⚠️ Skipping - No valid email address")
        continue
    
    # Find consolidated file for this ZBM
    safe_zbm_name = str(zbm_name).replace(' ', '_').replace('/', '_').replace('\\', '_')
    consolidated_pattern = os.path.join(consolidated_folder, f"ZBM_Consolidated_{zbm_code}_*.xlsx")
    consolidated_files = glob.glob(consolidated_pattern)
    
    if not consolidated_files:
        print(f"   ⚠️ No consolidated file found for {zbm_code}")
        print(f"      Searched pattern: {consolidated_pattern}")
        continue
    
    consolidated_file = consolidated_files[0]
    
    # Convert to absolute path (Outlook requires absolute paths)
    consolidated_file = os.path.abspath(consolidated_file)
    
    # Verify file exists
    if not os.path.exists(consolidated_file):
        print(f"   ❌ File does not exist: {consolidated_file}")
        continue
    
    print(f"   📎 Attaching: {os.path.basename(consolidated_file)}")
    print(f"      Full path: {consolidated_file}")
    
    # Read summary report data
    summary_html = read_summary_report(zbm_code, zbm_name)
    
    if not summary_html:
        print(f"   ⚠️ No summary report data found for {zbm_code}")
        continue
    
    # Get ABM emails for CC
    abm_cc_emails = get_abm_emails_for_zbm(zbm_code)
    
    # Create email
    try:
        mail = outlook.CreateItem(0)
        mail.To = zbm_email
        
        # Add ABMs in CC (uncomment if needed)
        # if abm_cc_emails:
        #     mail.CC = abm_cc_emails
        #     print(f"   📧 CC: {len(abm_cc_emails.split(';'))} ABMs")
        
        # Set subject
        mail.Subject = f"Sample Direct Dispatch - ZBM Summary Report as of {current_date}"
        
        # Render email body with summary table
        mail.HTMLBody = template.render(
            zbm_name=zbm_name,
            zbm_code=zbm_code,
            current_date=current_date,
            summary_table=summary_html
        )
        
        # Set sender
        mail.SentOnBehalfOfName = 'EPD_SFA@abbott.com'
        
        # Attach consolidated file - AFTER setting body
        try:
            mail.Attachments.Add(consolidated_file)
            print(f"   ✅ Attachment added successfully")
        except Exception as attach_error:
            print(f"   ❌ Error attaching file: {attach_error}")
            print(f"      File path: {consolidated_file}")
            continue
        
        # Display email
        mail.Display()
        
        email_count += 1
        print(f"   ✅ Email displayed successfully for {zbm_email}")
        
        # Log the sent email
        with open(os.path.join(email_log_folder, 'email_log.txt'), 'a') as log:
            log.write(f"{dt.now()} - Displayed email for {zbm_code} ({zbm_name}) - {zbm_email}\n")
        
    except Exception as e:
        print(f"   ❌ Error creating email for {zbm_code}: {e}")
        import traceback
        traceback.print_exc()
        continue

print(f"\n🎉 Email automation completed!")
print(f"📊 Total emails displayed: {email_count} out of {len(zbms)} ZBMs")
print(f"📁 Email logs saved in: {email_log_folder}")
The key changes I made to remove the "None" row:
Added proper filtering after creating the DataFrame to remove rows where the first column contains "None" (case-insensitive)
Added .reset_index(drop=True) to properly reset the DataFrame index after filtering
Fixed the filter chain with proper .str.strip().str.lower() to handle the "None" text
Added traceback for better error debugging
This will now completely remove the row with "None" in the first column from the email summary table while keeping the Total row intact.
